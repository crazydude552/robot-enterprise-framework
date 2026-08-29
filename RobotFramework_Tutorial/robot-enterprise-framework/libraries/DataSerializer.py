import openpyxl
import requests
from lxml import etree


class DataSerializer:
    ROBOT_LIBRARY_SCOPE = "GLOBAL"

    def read_excel_test_data(self, file_path, test_case_id):
        """Reads row from Excel matching TestCaseID."""
        wb = openpyxl.load_workbook(file_path)
        sheet = wb.active
        headers = [cell.value for cell in sheet[1]]

        for row in sheet.iter_rows(min_row=2, values_only=True):
            row_dict = dict(zip(headers, row))
            if (
                str(row_dict.get("TestCaseID")).strip()
                == str(test_case_id).strip()
            ):
                return row_dict

        raise ValueError(
            f"TestCaseID '{test_case_id}' not found in Excel: {file_path}"
        )

    def process_and_serialize_xml(self, template_path, combined_data):
        """Substitutes ${KEY} placeholders in the XML template."""
        with open(template_path, "r", encoding="utf-8") as f:
            content = f.read()

        for key, value in combined_data.items():
            placeholder = f"${{{key}}}"
            content = content.replace(
                placeholder, str(value) if value is not None else ""
            )

        return content

    def validate_xml_against_xsd(self, xml_string, xsd_path):
        """Validates XML payload against XSD schema."""
        with open(xsd_path, "rb") as xsd_file:
            schema_doc = etree.parse(xsd_file)
            schema = etree.XMLSchema(schema_doc)

        xml_doc = etree.fromstring(xml_string.encode("utf-8"))
        if not schema.validate(xml_doc):
            errors = "\n".join(
                [f"Line {e.line}: {e.message}" for e in schema.error_log]
            )
            raise ValueError(f"XSD Validation Failure:\n{errors}")

        return True

    def send_swift_api_request(self, url, xml_payload, user, password, timeout):
        """Sends HTTP POST request with Basic Authentication."""
        headers = {
            "Content-Type": "application/xml",
            "Accept": "application/xml",
        }
        res = requests.post(
            url,
            data=xml_payload.encode("utf-8"),
            headers=headers,
            auth=(user, password),
            timeout=float(timeout),
        )
        res.raise_for_status()
        return res.text

    def parse_pacs002_response(self, response_xml):
        """Extracts status and error reason codes from pacs.002 responses using XPath."""
        xml_doc = etree.fromstring(response_xml.encode("utf-8"))

        def get_text(xpath_expr):
            nodes = xml_doc.xpath(xpath_expr)
            return nodes[0].text.strip() if nodes and nodes[0].text else ""

        return {
            "TransactionStatus": get_text("//*[local-name()='TxSts']"),
            "ReasonCode": get_text(
                "//*[local-name()='Rsn']/*[local-name()='Cd']"
            ),
            "AdditionalInfo": get_text("//*[local-name()='AddtlInf']"),
            "OriginalUETR": get_text("//*[local-name()='OrgnlUETR']"),
        }

    def parse_camt053_statement(self, xml_string):
        """Extracts account balances and entries from a camt.053 XML string."""
        xml_doc = etree.fromstring(xml_string.encode("utf-8"))

        def get_text(xpath_expr):
            nodes = xml_doc.xpath(xpath_expr)
            return nodes[0].text.strip() if nodes and nodes[0].text else ""

        opening_bal = get_text(
            "//*[local-name()='Bal'][*[local-name()='Tp']/*[local-name()='CdOrPrtry']/*[local-name()='Cd']='OPBD']/*[local-name()='Amt']"
        )
        closing_bal = get_text(
            "//*[local-name()='Bal'][*[local-name()='Tp']/*[local-name()='CdOrPrtry']/*[local-name()='Cd']='CLBD']/*[local-name()='Amt']"
        )
        entry_amt = get_text(
            "//*[local-name()='Ntry']/*[local-name()='Amt']"
        )
        uetr = get_text("//*[local-name()='UETR']")

        return {
            "OpeningBalance": opening_bal,
            "ClosingBalance": closing_bal,
            "EntryAmount": entry_amt,
            "UETR": uetr,
        }